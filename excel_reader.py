"""Read administrative workbooks in memory, retaining source-row provenance."""
from dataclasses import dataclass, field
from io import BytesIO
import re
import warnings
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import MergedCell

from normalizer import build_rib, name, text, rib
from row_classifier import classify_rows, build_regions, RowClassification, SUMMARY_TYPES, AMOUNT_HEADERS, parse_summary_amount, summary_label


class WorkbookError(ValueError):
    pass


@dataclass
class Record:
    row: int
    values: dict
    issues: dict = field(default_factory=dict)


@dataclass
class Table:
    sheet: str
    header: int
    end_header: int
    columns: dict
    records: list = field(default_factory=list)
    all_columns: dict = field(default_factory=dict)
    cells: dict = field(default_factory=dict)
    rows: list = field(default_factory=list)
    detection: str = "En-têtes administratifs reconnus"
    original_rows: list = field(default_factory=list)
    row_classifications: dict = field(default_factory=dict)
    row_schemas: dict = field(default_factory=dict)
    regions: list = field(default_factory=list)
    merged_ranges: list = field(default_factory=list)

    @property
    def available(self):
        cols = self.columns
        return {f for f, ok in {
            "name": "name" in cols or {"first", "last"} <= cols.keys(),
            "ddp": "ddp" in cols, "cin": "cin" in cols,
            "rib": "rib" in cols or {"bank", "city", "account", "key"} <= cols.keys(),
        }.items() if ok}


@dataclass
class SourceCell:
    value: object
    display: str
    coordinate: str
    data_type: str
    number_format: str
    role: str = "DATA_VALUE"


def complete_columns(sheet, table):
    """Keep physical columns, duplicate headings and unnamed populated columns."""
    columns = {}
    for col in range(1, sheet.max_column + 1):
        labels = []
        for row in range(table.header, table.end_header + 1):
            value = text(cell_at(sheet, row, col).value).strip()
            if value and value not in labels:
                labels.append(value)
        if labels or any(sheet.cell(row, col).value is not None for row in range(1, sheet.max_row + 1)):
            columns[col] = " / ".join(labels) or f"Sans titre ({get_column_letter(col)})"
    table.all_columns = columns
    return table


def header_key(value):
    # Some source documents literally contain the four characters \\xb0.
    return re.sub(r"[^A-Z0-9]", "", name(value).replace("\\XB0", ""))


ALIASES = {
    "name": ["NOM ET PRENOM", "NOM PRENOM", "NOM ET PRENOMS"],
    "last": ["NOM"], "first": ["PRENOM", "PRENOMS"],
    "ddp": ["N DDP", "DDP", "DDR", "N DDR", "NUMERO DDP"],
    "cin": ["CIN", "N CIN", "NUMERO CIN"], "rib": ["RIB", "R.I.B"],
    "bank": ["BANQUE", "CODE BANQUE"], "city": ["VILLE", "CODE VILLE"],
    "account": ["NUMERO COMPTE RIB", "NUMORO COMPTE RIB", "NUMEORO COMPTE RIB", "NUMEORO COMPTE RIBE", "COMPTE RIB", "NUMERO DE COMPTE RIB"],
    "key": ["CLE", "CLE RIB"],
}
LOOKUP = {header_key(alias): field for field, aliases in ALIASES.items() for alias in aliases}


def open_workbook(data):
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Cannot parse header or footer.*")
            return load_workbook(BytesIO(data), data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, TypeError) as exc:
        raise WorkbookError("Impossible de lire ce fichier. Choisissez un classeur .xlsx valide, non protégé par mot de passe.") from exc


def cell_at(sheet, row, column):
    cell = sheet.cell(row, column)
    if cell.value is None:
        for merged in sheet.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
                return sheet.cell(merged.min_row, merged.min_col)
    return cell


def detect_header(sheet, row):
    columns = {}
    end = row
    for col in range(1, sheet.max_column + 1):
        cell = cell_at(sheet, row, col)
        key = LOOKUP.get(header_key(cell.value))
        if key and key not in columns:
            columns[key] = cell.column
            for merged in sheet.merged_cells.ranges:
                if cell.coordinate in merged:
                    end = max(end, merged.max_row)
    return Table(sheet.title, row, end, columns)


def detect_tables(workbook):
    tables = []
    for sheet in workbook:
        candidates = [detect_header(sheet, row) for row in range(1, sheet.max_row + 1)]
        candidates = [t for t in candidates if len(t.available) >= 2]
        if candidates:
            # Select the first header block, not a richer later table that would
            # hide the records preceding it. Keep multi-level header support.
            block = [candidates[0]]
            for candidate in candidates[1:]:
                if candidate.header > max(t.end_header for t in block) + 1:
                    break
                block.append(candidate)
            tables.append(max(block, key=lambda t: (len(t.available), len(t.columns), -t.header)))
    return tables


def discover_tables(workbook):
    """Inventory every nonempty sheet; generic guesses are explicitly reviewable."""
    known = {t.sheet: t for t in detect_tables(workbook)}
    result = []
    for sheet in workbook:
        if sheet.title in known:
            result.append(known[sheet.title])
            continue
        populated = [(row[0].row, [c for c in row if c.value is not None])
                     for row in sheet.iter_rows()]
        populated = [(r, cells) for r, cells in populated if cells]
        if not populated:
            continue
        # Only structure is inferred; unknown meanings and business rules are not.
        candidates = [(r, cells) for r, cells in populated if len(cells) >= 2 and
                      all(isinstance(c.value, str) and c.data_type != "f" for c in cells)]
        row = candidates[0][0] if candidates else populated[0][0]
        t = detect_header(sheet, row)
        t.detection = "En-têtes proposés par la structure : à vérifier"
        result.append(t)
    return result


def read_cell(cell):
    value = cell.value
    if cell.data_type == "f":
        return text(value), "Formule : remplacez-la par sa valeur dans Excel."
    if cell.data_type == "e":
        return text(value), "Erreur Excel : corrigez la cellule source."
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if abs(value) >= 10 ** 15:
            return text(value), "Précision incertaine : identifiant numérique de plus de 15 chiffres. Vérifiez la source et enregistrez-le comme texte."
        if isinstance(value, float) and not value.is_integer():
            return text(value), "Identifiant numérique contenant des décimales."
        if re.fullmatch(r"0+", cell.number_format):
            return text(value).zfill(len(cell.number_format)), ""
    return text(value), ""


def read_table(workbook, table, overrides=None):
    sheet = workbook[table.sheet]
    complete_columns(sheet, table)
    # Find later header blocks, including tables with reordered columns.
    schemas, active = {}, table
    for row in range(1, sheet.max_row + 1):
        if row > active.end_header:
            candidate = detect_header(sheet, row)
            repeated = sum(bool(sheet.cell(row, col).value) and
                           header_key(sheet.cell(row, col).value) == header_key(title)
                           for col, title in active.all_columns.items()) >= 2
            if len(candidate.available) >= 2 or repeated:
                active = complete_columns(sheet, candidate)
        schemas[row] = active
    classifications = classify_rows(sheet, table, LOOKUP, header_key, schemas)
    for row, kind in (overrides or {}).items():
        if row not in classifications or kind not in {"DATA", "HEADER", "TITLE", "TOTAL", "SUBTOTAL", "REPORT", "SUMMARY", "EMPTY", "NOTES", "UNKNOWN", "OTHER"}:
            raise WorkbookError("Correction de structure invalide : vérifiez les lignes et le type choisi.")
        if kind == "EMPTY" and any(sheet.cell(row, col).value is not None for col in table.all_columns):
            raise WorkbookError("Une ligne contenant des valeurs ne peut pas être déclarée vide.")
        amounts = []
        if kind in SUMMARY_TYPES:
            for col in table.all_columns:
                try:
                    parse_summary_amount(sheet.cell(row, col).value)
                    amounts.append(col)
                except ValueError:
                    cell = sheet.cell(row, col)
                    if cell.value is not None and header_key(schemas[row].all_columns[col]) in AMOUNT_HEADERS and not summary_label(cell.value):
                        amounts.append(col)
        classifications[row] = RowClassification(kind, "Classification corrigée manuellement dans les options avancées.", tuple(amounts), 1.0)
    records = []
    cells, rows = {}, []
    for row in range(1, sheet.max_row + 1):
        classification = classifications[row]
        for col in table.all_columns:
            cell = sheet.cell(row, col)  # Never forward-fill merged data cells.
            role = "MERGED_PLACEHOLDER" if isinstance(cell, MergedCell) else "DATA_VALUE" if classification.kind == "DATA" else "SUMMARY_AMOUNT" if col in classification.amount_columns else "STRUCTURAL"
            if role == "DATA_VALUE":
                title = header_key(schemas[row].all_columns[col])
                canonical = LOOKUP.get(title)
                if canonical in {"cin", "ddp", "bank", "city", "account", "key"}:
                    role = "IDENTIFIER"
                elif canonical == "rib":
                    role = "BANK_IDENTIFIER"
                elif canonical in {"name", "first", "last"}:
                    role = "TEXT"
                elif title in {"DEBUT", "FIN", "DATEDEBUT", "DATEFIN", "DATENAISSANCE"} or title.startswith("DATE"):
                    role = "DATE"
                elif title in {"NETDH", "MONTANT", "NETAPAYER", "MONTANTNET", "BRUT", "RETENUES"}:
                    role = "AMOUNT"
            elif role == "STRUCTURAL" and classification.kind in SUMMARY_TYPES and cell.value is not None:
                role = "SUMMARY_TEXT"
            cells[row, col] = SourceCell(cell.value, read_cell(cell)[0], cell.coordinate, cell.data_type, cell.number_format, role)
        if classification.kind != "DATA":
            continue
        rows.append(row)
        raw, errors = {}, {}
        for key, col in schemas[row].columns.items():
            raw[key], error = read_cell(sheet.cell(row, col))
            if error:
                errors[key] = error
        values = {
            "name": raw.get("name", " ".join(raw.get(k, "") for k in ("first", "last"))),
            "ddp": raw.get("ddp", ""), "cin": raw.get("cin", ""),
            "rib": raw.get("rib", build_rib(*(raw.get(k, "") for k in ("bank", "city", "account", "key")))),
        }
        issues = {}
        dependencies = {"name": ["name"] if "name" in raw else ["first", "last"],
                        "ddp": ["ddp"], "cin": ["cin"],
                        "rib": ["rib"] if "rib" in raw else ["bank", "city", "account", "key"]}
        for key, sources in dependencies.items():
            messages = [errors[s] for s in sources if s in errors]
            if any(not raw.get(s, "").strip() for s in sources):
                messages.append("Information manquante.")
            if messages:
                issues[key] = " ".join(dict.fromkeys(messages))
        if "rib" not in raw:
            for key, width in zip(("bank", "city", "account", "key"), (3, 3, 16, 2)):
                part = rib(raw.get(key, ""))
                if part and (not part.isascii() or not part.isdigit() or len(part) > width):
                    issues["rib"] = "Composant RIB invalide (longueurs attendues : 3 / 3 / 16 / 2)."
        full = rib(values["rib"])
        if full and (not full.isascii() or not full.isdigit() or len(full) != 24):
            issues["rib"] = "RIB invalide : 24 chiffres attendus. " + issues.get("rib", "")
        records.append(Record(row, values, issues))
    return Table(table.sheet, table.header, table.end_header, table.columns, records,
                 table.all_columns, cells, rows, table.detection, list(range(1, sheet.max_row + 1)),
                 classifications, schemas, build_regions(classifications, schemas), list(map(str, sheet.merged_cells.ranges)))
