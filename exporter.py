"""Create safe downloadable workbooks in memory."""
from io import BytesIO
from copy import copy, deepcopy

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.cell.cell import MergedCell

from comparator import STATUSES


SHEETS = {"MATCH": "Correspondances", "DIFFERENT": "Différences", "MISSING_IN_B": "Absents",
          "INCOMPLETE": "Incomplets", "AMBIGUOUS": "Correspondances ambiguës"}


def write_rows(sheet, rows):
    for row in rows:
        sheet.append(list(row))
        for cell in sheet[sheet.max_row]:
            # Excel must never execute uploaded text as a formula.
            if isinstance(cell.value, str):
                cell.data_type = "s"
                cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="176B67")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for i, column in enumerate(sheet.columns, 1):
        width = max((len(str(c.value or "")) for c in column), default=12)
        sheet.column_dimensions[get_column_letter(i)].width = min(55, max(16, width + 2))


def export_excel(results, metadata=None):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Résumé"
    counts = results["Code statut"].value_counts() if not results.empty else {}
    rows = [["Indicateur", "Valeur"], ["Total analysé", len(results)]]
    rows.extend([STATUSES[key], int(counts.get(key, 0))] for key in STATUSES)
    rows.extend([key, str(value)] for key, value in (metadata or {}).items())
    write_rows(summary, rows)
    for status, title in SHEETS.items():
        sheet = workbook.create_sheet(title)
        if results.empty:
            write_rows(sheet, [["Statut"]])
        else:
            subset = results.loc[results["Code statut"] == status].drop(columns="Code statut")
            write_rows(sheet, [subset.columns.tolist(), *subset.itertuples(index=False, name=None)])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


COLORS = {"Erreur": "FADBD8", "Avertissement": "FFF0C2", "Information": "DDEBF7"}


def export_quality(report, workbook_a, workbook_b, tables_a, tables_b, a, b, mappings, rules_a, rules_b):
    """Preserve source cell values, styles and coordinates in separate annotated copies.

    Formula expressions are retained as inert text; never execute external source formulas.
    """
    from openpyxl import load_workbook
    from quality_report import errors_only_report
    import pandas as pd
    from data_quality import FINDING_COLUMNS
    report = errors_only_report(report)
    visible = report.get("visible_findings")
    annotations = report["findings"] if visible is None else visible
    anomalies = report["anomalies"] if visible is None else pd.concat(
        [report["anomalies"], visible[FINDING_COLUMNS]], ignore_index=True).drop_duplicates()
    workbook = load_workbook(BytesIO(export_excel(report["records"])))
    summary = workbook["Résumé"]
    summary.title = "Résumé global"
    # Rebuild to keep a single consistent table and include units and analysis coverage.
    summary.delete_rows(1, summary.max_row)
    rows = [["Indicateur", "Valeur"], ["Lignes A analysées (rapprochement)", len(report["records"])]]
    rows.extend([label, int((report["records"]["Code statut"] == status).sum()) if not report["records"].empty else 0] for status, label in STATUSES.items())
    rows.extend(report["stats"].items())
    rows.extend(report.get("main_metrics", {}).items())
    rows.extend([["Unité des anomalies", "Une règle signalée dans une cellule ; plusieurs règles peuvent concerner la même cellule."],
                 ["Unité des différences", "Une paire de valeurs par ligne rapprochée et champ comparé ; aucune source n'est présumée correcte."],
                 ["Feuilles rapprochées", f"A : {a.sheet} / B : {b.sheet}"],
                 ["Formules", "Conservées comme texte non exécutable, sans recalcul ; les valeurs sources ne sont pas corrigées."],
                 ["Couleurs", "Rouge : erreur ou différence confirmée uniquement."]])
    write_rows(summary, rows)
    frames = {
        "Anomalies fichier A": anomalies.loc[anomalies["Fichier"] == "A"],
        "Anomalies fichier B": anomalies.loc[anomalies["Fichier"] == "B"],
        "Différences entre fichiers": report["differences"],
        "Colonnes non comparées": report["unmatched"],
        "Correspondances colonnes": mappings,
        "Règles A": rules_a, "Règles B": rules_b,
    }
    if "missing_people" in report:
        absent = report["missing_people"]
        frames["Dans A absentes de B"] = absent.loc[absent["Fichier"] == "A"]
        frames["Dans B absentes de A"] = absent.loc[absent["Fichier"] == "B"]
    for title, frame in frames.items():
        clean = frame.astype(object).where(frame.notna(), "")
        write_rows(workbook.create_sheet(title), [clean.columns.tolist(), *clean.itertuples(index=False, name=None)])
    inventory = [["Fichier", "Feuille source", "Feuille annotée", "En-têtes", "Lignes", "Colonnes"]]
    structures = [["Fichier", "Feuille", "Ligne Excel", "Type", "Confiance", "Raison"]]
    regions = [["Fichier", "Feuille", "Début", "Fin", "Type", "En-têtes", "Confiance", "Raison"]]
    roles = [["Fichier", "Feuille", "Cellule", "Rôle", "Type de ligne"]]
    severity_order = {key: i for i, key in enumerate(COLORS)}
    for label, source_book, tables, main in (("A", workbook_a, tables_a, a), ("B", workbook_b, tables_b, b)):
        ordered = sorted(tables, key=lambda t: t.sheet != main.sheet)
        for index, table in enumerate(ordered):
            structures.extend([label, table.sheet, row, info.kind, info.confidence, info.reason]
                              for row, info in table.row_classifications.items())
            regions.extend([label, table.sheet, *region.values()] for region in table.regions)
            roles.extend([label, table.sheet, cell.coordinate, cell.role, table.row_classifications[row].kind]
                         for (row, col), cell in table.cells.items())
            source = source_book[table.sheet]
            title = f"Fichier {label} annoté" + (f" {index + 1}" if index else "")
            target = workbook.create_sheet(title)
            for row in source:
                for cell in row:
                    dest = target.cell(cell.row, cell.column, cell.value)
                    if cell.has_style:
                        # Style IDs are workbook-local; copy actual style objects.
                        dest.font = copy(cell.font)
                        dest.fill = copy(cell.fill)
                        dest.border = copy(cell.border)
                        dest.alignment = copy(cell.alignment)
                        dest.protection = copy(cell.protection)
                        dest.number_format = cell.number_format
                    if cell.data_type == "f" or isinstance(cell.value, str):
                        dest.data_type = "s"
                    if cell.comment:
                        dest.comment = copy(cell.comment)
            for key, dimension in source.column_dimensions.items():
                target.column_dimensions[key] = copy(dimension)
            for key, dimension in source.row_dimensions.items():
                target.row_dimensions[key] = copy(dimension)
            for merged in source.merged_cells.ranges:
                target.merge_cells(str(merged))
            target.freeze_panes = source.freeze_panes or f"A{table.end_header + 1}"
            target.sheet_format = copy(source.sheet_format)
            target.sheet_properties = copy(source.sheet_properties)
            target.auto_filter = deepcopy(source.auto_filter)
            marks = annotations.loc[(annotations["Fichier"] == label) & (annotations["Feuille"] == table.sheet)]
            for coordinate, group in marks.groupby("Cellule"):
                cell = target[coordinate]
                if isinstance(cell, MergedCell):
                    continue
                severity = min(group["Gravité"], key=severity_order.get)
                cell.fill = PatternFill("solid", fgColor=COLORS[severity])
                explanations = "\n".join(dict.fromkeys(f"{r['Gravité']} — {r['Type']} : {r['Explication']}" for r in group.to_dict("records")))
                if cell.comment:
                    explanations = cell.comment.text + "\n\n" + explanations
                cell.comment = Comment(explanations, "Contrôle Excel")
            inventory.append([label, table.sheet, title, table.header, len(table.rows), len(table.all_columns)])
    write_rows(workbook.create_sheet("Périmètre analysé"), inventory)
    write_rows(workbook.create_sheet("Classification des lignes"), structures)
    write_rows(workbook.create_sheet("Régions détectées"), regions)
    write_rows(workbook.create_sheet("Rôles des cellules"), roles)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
