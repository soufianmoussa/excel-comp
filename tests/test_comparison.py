from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from comparator import compare
from excel_reader import Record, detect_tables, open_workbook, read_cell, read_table, WorkbookError
from exporter import export_excel
from normalizer import build_rib, cin, identifier, name


# Standard RIB component widths are 3/3/16/2 (24 total).
RIB = build_rib("350", "810", "0000000000459889", "40")


def rec(row=2, **changes):
    values = {"name": "SAID EXEMPLE", "ddp": "721155", "cin": "X123456", "rib": RIB}
    values.update(changes)
    return Record(row, values)


def statuses(a, b, fields=("name", "ddp", "cin", "rib")):
    return compare(a, b, fields)["Code statut"].tolist()


def test_normalizers():
    assert name("  Saïd\u00a0  Exemple ") == "SAID EXEMPLE"
    assert name(None) == ""
    assert identifier(721155.0) == identifier("721155")
    assert identifier("00721155.0") == "00721155"
    assert cin(" x 123456 ") == "X123456"
    assert len(RIB) == 24
    assert RIB == "350" + "810" + "0000000000459889" + "40"
    assert build_rib(1, 2, 3, 4) == "001002000000000000000304"
    assert build_rib(1, 2, None, 4) == ""


def test_reordered():
    other = rec(3, name="AUTRE PERSONNE", ddp="2", cin="Y2", rib=build_rib(1, 2, 3, 4))
    assert statuses([rec(), other], [other, rec()]) == ["MATCH", "MATCH"]


def test_missing_and_different():
    assert statuses([rec()], []) == ["MISSING_IN_B"]
    result = compare([rec()], [rec(rib=RIB[:-1] + "1")], ["name", "ddp", "cin", "rib"])
    assert result.iloc[0]["Code statut"] == "DIFFERENT"
    assert result.iloc[0]["Champs différents"] == "RIB"
    assert result.iloc[0]["RIB fichier A"] == RIB


def test_duplicates_and_conflicting_identifiers():
    assert statuses([rec()], [rec(), rec(3)]) == ["AMBIGUOUS"]
    assert statuses([rec(), rec(3)], [rec()]) == ["MATCH", "MATCH"]
    assert statuses([rec()], [rec(cin="OTHER"), rec(3, ddp="OTHER", rib="OTHER")]) == ["AMBIGUOUS"]


def test_missing_values_are_never_matching_keys():
    assert statuses([rec(cin="")], [rec(cin="")]) == ["INCOMPLETE"]
    assert statuses([rec(cin="")], [], ["cin"]) == ["INCOMPLETE"]
    assert statuses([rec(cin="")], [rec(cin="")], ["ddp"]) == ["MATCH"]
    with pytest.raises(ValueError):
        compare([rec()], [rec()], [])


def test_name_fallback_and_selected_fields():
    assert statuses([rec()], [rec(cin="Z", ddp="Z", rib="Z")]) == ["DIFFERENT"]
    assert statuses([rec()], [rec(cin="Z")], ["cin"]) == ["MISSING_IN_B"]
    # Same names alone must not create ambiguity when IDs disambiguate.
    assert statuses([rec()], [rec(), rec(3, ddp="Z", cin="Z", rib="Z")]) == ["MATCH"]


def administrative_book():
    w = Workbook()
    s = w.active
    s.title = "Tableau"
    s["A1"] = "Document administratif"
    s.merge_cells("D3:D5")
    s["D3"] = "R.I.B"
    for col, header in zip("ABC", ["NOM ET PRENOM", "N° DDP", "N° CIN"]):
        s.merge_cells(f"{col}4:{col}5")
        s[f"{col}4"] = header
    for col, value in enumerate(["SAID EXEMPLE", 721155, "X123456", RIB], 1):
        s.cell(6, col, value)
    w.create_sheet("Vide")
    return w


def test_multilevel_merged_header():
    w = administrative_book()
    tables = detect_tables(w)
    assert len(tables) == 1
    table = read_table(w, tables[0])
    assert table.header == 4 and table.end_header == 5
    assert table.available == {"name", "ddp", "cin", "rib"}
    assert len(table.records) == 1
    assert table.records[0].values["rib"] == RIB


def test_split_columns_and_multiple_sheets():
    w = Workbook()
    s = w.active
    s.append(["CIN", "Nom", "Prénom", "DDR", "Banque", "Ville", "Numéoro compte Rib", "Clé"])
    s.append(["X123456", "EXEMPLE", "SAID", 721155, 350, 810, "0000000000459889", 40])
    w.copy_worksheet(s)
    assert len(detect_tables(w)) == 2
    r = read_table(w, detect_tables(w)[0]).records[0]
    assert r.values["rib"] == RIB and r.values["name"] == "SAID EXEMPLE"
    assert not r.issues


def test_numeric_precision_and_formats():
    w = administrative_book()
    s = w.active
    s["D6"] = 123456789012345678901234
    r = read_table(w, detect_tables(w)[0]).records[0]
    assert "Précision incertaine" in r.issues["rib"]
    assert statuses([r], [rec()]) == ["INCOMPLETE"]
    s["B6"] = 123
    s["B6"].number_format = "000000"
    assert read_cell(s["B6"])[0] == "000123"
    s["B6"] = "=1+1"
    assert "Formule" in read_cell(s["B6"])[1]


def test_export_text_and_formula_safety():
    result = compare([rec(name="=1+1")], [rec(name="=1+1")], ["cin"])
    data = export_excel(result)
    w = load_workbook(BytesIO(data), data_only=False)
    assert w.sheetnames == ["Résumé", "Correspondances", "Différences", "Absents", "Incomplets", "Correspondances ambiguës"]
    cells = list(w["Correspondances"].iter_rows(min_row=2))[0]
    assert all(c.data_type != "f" for c in cells)
    assert any(c.value == RIB and c.data_type == "s" for c in cells)
    assert any(c.value == "=1+1" and c.data_type == "s" for c in cells)


def test_corrupt_file():
    with pytest.raises(WorkbookError):
        open_workbook(b"not an Excel file")


def test_streamlit_workflow(monkeypatch):
    import streamlit as st
    from streamlit.testing.v1 import AppTest

    app_path = Path(__file__).resolve().parents[1] / "app.py"
    initial = AppTest.from_file(str(app_path)).run()
    assert not initial.exception
    assert initial.button[0].disabled
    data = BytesIO()
    administrative_book().save(data)
    monkeypatch.setattr(st, "file_uploader", lambda *args, **kwargs: BytesIO(data.getvalue()))
    app = AppTest.from_file(str(app_path)).run()
    assert not app.exception
    assert not app.button[0].disabled
    app.button[0].click().run()
    assert not app.exception
    assert app.session_state["results"]["Code statut"].tolist() == ["MATCH"]
    assert len(app.tabs) == 8  # Two directional absence lists plus existing details.
    assert len(app.session_state["export"]) > 0
    assert len(app.metric) == 3
    assert not any(str(c.key).startswith("field_") for c in app.checkbox)
    assert app.button[0].label == "Analyser et comparer"


def test_local_samples_structure_only():
    root = Path(__file__).resolve().parents[1]
    for filename in ("FILE_A.xlsx", "File_B.xlsx"):
        path = root / filename
        if not path.exists():
            pytest.skip("Exemples locaux non fournis")
        w = open_workbook(path.read_bytes())
        tables = detect_tables(w)
        assert len(tables) == 1
        table = read_table(w, tables[0])
        assert table.available == {"name", "ddp", "cin", "rib"}
        # Local examples can grow; synthetic fixtures assert exact record counts.
        assert len(table.records) >= 2
        assert len(table.rows) == len(table.records)
        w.close()
