from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from data_quality import Rule, analyze_table, rules_frame
from excel_reader import discover_tables, read_table
from exporter import export_quality
from field_mapping import detect_mappings
from normalizer import FIELDS
from quality_report import analyze, errors_only_report, presentation_summary
from quality_ui import table_frame, styled_table
from row_classifier import parse_summary_amount


HEADERS = ["DEBUT", "CIN", "DDR", "NOM ET PRENOM", "NET(DH)"]


def load(book, overrides=None):
    return read_table(book, discover_tables(book)[0], overrides)


def administrative():
    w = Workbook()
    s = w.active
    s.append(["État au 01/10/2025"])
    s.merge_cells("A1:E1")
    s.append(HEADERS)
    s.append(["31/02/2025", "X1", 1, "Personne Exemple", "invalide"])
    s.append(["Report........................", None, None, None, "15.147,60 dhs"])
    s.merge_cells("A4:D4")
    s.append(["01/10/2025", "X2", 2, "Autre Exemple", 10])
    s.append(["Total général", None, None, None, "15.157,60 dhs"])
    s.merge_cells("A6:D6")
    s.append(["Sous-total", None, None, None, 10])
    s.merge_cells("A7:D7")
    s.append(["01/10/2025", None, 3, "Personne Partielle", 10])
    s.append(["Séparation inhabituelle"])
    s.merge_cells("A9:E9")
    s["A9"].font = Font(bold=True)
    s.append(["01/10/2025", "X4", 4, "Total Martin", 10])
    s.append(["Bilan intermédiaire", None, None, None, 30])
    s.merge_cells("A11:D11")
    s["A11"].font = Font(bold=True)
    s.append(["01/10/2025", "X5", 5, "Dernier Exemple", 10])
    s.append(["Note : fin du groupe"])
    s.merge_cells("A13:E13")
    s.append([None] * 5)
    return w


def report(a, b):
    m = detect_mappings(a, b, list(FIELDS))
    return analyze([a], [b], a, b, list(FIELDS), m, {}, {}), m


def test_roles_preserve_all_cells_and_true_errors():
    w = administrative()
    t = load(w)
    expected = {1: "TITLE", 2: "HEADER", 3: "DATA", 4: "REPORT", 5: "DATA",
                6: "TOTAL", 7: "SUBTOTAL", 8: "DATA", 9: "UNKNOWN", 10: "DATA",
                11: "SUMMARY", 12: "DATA", 13: "NOTES", 14: "EMPTY"}
    assert {r: info.kind for r, info in t.row_classifications.items()} == expected
    assert t.rows == [3, 5, 8, 10, 12]
    assert len(t.cells) == 14 * 5
    assert t.cells[4, 1].value == w.active["A4"].value
    assert t.cells[4, 2].value is None
    assert t.cells[4, 2].coordinate == "B4"
    assert t.cells[3, 1].role == "DATE"
    assert t.cells[3, 2].role == "IDENTIFIER"
    assert t.cells[3, 5].role == "AMOUNT"
    assert t.cells[4, 5].role == "SUMMARY_AMOUNT"
    findings = analyze_table(t, "A")
    errors = findings.loc[findings["Gravité"] == "Erreur"]
    assert {"A3", "E3", "B8"} <= set(errors["Cellule"])
    assert set(errors["Ligne"]) <= {3, 8}
    assert all("DATA" in value for value in errors["Contexte de validation"])
    assert {r["Type"] for r in t.regions} >= {"REPORT_REGION", "SUMMARY_REGION", "UNKNOWN_REGION", "NOTES_REGION", "DATA_REGION"}
    assert t.row_classifications[9].confidence < .5
    assert list(table_frame(t).index) == list(range(1, 15))
    assert table_frame(t).at[4, "Type de ligne"] == "Report"
    styled = styled_table(t, [4], findings)
    assert not styled._compute().ctx


@pytest.mark.parametrize("row", [4, 22, 203])
@pytest.mark.parametrize("title", ["DEBUT", "NET(DH)"])
def test_embedded_merged_report_independent_of_row_or_column_type(row, title):
    w = Workbook()
    s = w.active
    s.append([title, "CIN", "DDR", "Note"])
    s.append(["incorrect", "X1", 1, None])
    s.cell(row, 1, "Report... 12.274,00 dhs")
    s.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    s.cell(row + 1, 2, "X2")
    s.cell(row + 1, 3, 2)
    t = load(w)
    assert t.row_classifications[row].kind == "REPORT"
    assert t.rows == [2, row + 1]
    errors = analyze_table(t, "B")
    assert "A2" in set(errors["Cellule"])
    assert row not in set(errors["Ligne"])
    assert t.cells[row, 1].value == "Report... 12.274,00 dhs"


def test_repeated_and_reordered_headers_compare_actual_cells():
    wa = Workbook()
    s = wa.active
    s.append(HEADERS)
    s.append(["01/10/2025", "X1", 1, "Premier", 10])
    s.append(["Report", None, None, None, 10])
    s.append(["CIN", "NET(DH)", "NOM ET PRENOM", "DDR", "DEBUT"])
    s.append(["X2", 20, "Second", 2, "02/10/2025"])
    wb = Workbook()
    wb.active.append(HEADERS)
    wb.active.append(["02/10/2025", "X2", 2, "Second", 21])
    wb.active.append(["01/10/2025", "X1", 1, "Premier", 10])
    a, b = load(wa), load(wb)
    assert a.row_classifications[4].kind == "HEADER"
    assert a.records[1].values["cin"] == "X2"
    r, _ = report(a, b)
    errors = r["anomalies"].loc[r["anomalies"]["Gravité"] == "Erreur"]
    assert errors.empty
    differences = r["differences"].loc[r["differences"]["Gravité"] == "Erreur"]
    assert len(differences) == 1
    assert differences.iloc[0]["Cellules A"] == "B5"
    assert differences.iloc[0]["Cellules B"] == "E2"
    assert set(r["records"]["Ligne A"]) == {2, 5}


def test_manual_classification_revalidates_without_discarding_values():
    w = administrative()
    t = load(w)
    assert 9 not in set(analyze_table(t, "A")["Ligne"])
    corrected = load(w, {9: "DATA"})
    findings = analyze_table(corrected, "A")
    assert "A9" in set(findings.loc[findings["Gravité"] == "Erreur", "Cellule"])
    assert corrected.cells[9, 1].value == t.cells[9, 1].value
    assert corrected.row_classifications[9].confidence == 1
    with pytest.raises(ValueError):
        load(w, {9: "EMPTY"})


def test_summary_amounts_separate_rules_and_bad_numeric_data():
    w = administrative()
    t = load(w)
    findings = analyze_table(t, "A", {5: Rule(kind="number", maximum="5")})
    assert 4 not in set(findings["Ligne"])
    own_rules = analyze_table(t, "A", summary_rules={5: Rule(kind="number", maximum="100")})
    assert 4 in set(own_rules["Ligne"])
    assert parse_summary_amount("15.147,60 dhs") == parse_summary_amount("15147.60")
    w.active["E4"] = "15.14,60 dhs"
    invalid = analyze_table(load(w), "A")
    assert "Montant de synthèse invalide" in set(invalid.loc[invalid["Ligne"] == 4, "Type"])


def test_export_and_metrics_keep_structures_without_people_matches():
    w = administrative()
    a, b = load(w), load(administrative())
    r, m = report(a, b)
    r = errors_only_report(r)
    visible, metrics = presentation_summary(r, [a], [b], a, b)
    assert metrics["Nombre de lignes analysées"] == 10
    assert 4 not in set(r["records"]["Ligne A"])
    result = load_workbook(BytesIO(export_quality(r, w, w, [a], [b], a, b, m, rules_frame([a]), rules_frame([b]))))
    copied = result["Fichier A annoté"]
    assert copied["A4"].value == w.active["A4"].value
    assert copied["E4"].value == w.active["E4"].value
    assert "A4:D4" in copied.merged_cells
    assert copied["A4"].comment is None
    assert copied["A3"].comment is not None
    assert {"Classification des lignes", "Régions détectées", "Rôles des cellules"} <= set(result.sheetnames)


def test_only_structure_and_numeric_amount_in_identifier_position():
    w = Workbook()
    s = w.active
    s.append(["CIN", "DDR", "Note"])
    s.append(["Report", 12345, None])
    t = load(w)
    assert t.row_classifications[2].kind == "REPORT"
    r, _ = report(t, t)
    assert r["records"].empty
    assert "Code statut" in r["records"]
    assert r["anomalies"].empty


def test_first_table_kept_when_later_table_has_more_identifiers():
    w = Workbook()
    w.active.append(["CIN", "DDR"])
    w.active.append(["X1", 1])
    w.active.append(["CIN", "DDR", "NOM ET PRENOM"])
    w.active.append(["X2", 2, "Exemple"])
    t = load(w)
    assert t.header == 1 and t.rows == [2, 4]
    assert [r.values["cin"] for r in t.records] == ["X1", "X2"]


def test_headers_after_fifty_rows_and_generic_repeated_headers():
    w = Workbook()
    w.active.cell(65, 1, "CIN")
    w.active.cell(65, 2, "DDR")
    w.active.cell(66, 1, "X1")
    w.active.cell(66, 2, 1)
    assert load(w).rows == [66]
    w = Workbook()
    for row in [["Référence", "Commentaire"], [1, "Texte"], ["Référence", "Commentaire"], [2, "Suite"]]:
        w.active.append(row)
    t = load(w)
    assert t.rows == [2, 4]
    assert t.row_classifications[3].kind == "HEADER"


def test_manual_zone_ui_invalidates_report_and_applies_validation(monkeypatch):
    import streamlit as st
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    w = administrative()
    data = BytesIO()
    w.save(data)
    monkeypatch.setattr(st, "file_uploader", lambda *args, **kwargs: BytesIO(data.getvalue()))
    app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py")).run(timeout=15)
    app.button[0].click().run(timeout=15)
    assert not app.exception and not app.error
    assert any(e.label == "Structures à vérifier" for e in app.expander)
    assert 9 not in set(app.session_state["report"]["anomalies"]["Ligne"])
    key = next(k for k in app.session_state.filtered_state if k.startswith("zones_A_"))
    delta = {"edited_rows": {}, "added_rows": [{"Feuille": w.active.title, "Début": 9, "Fin": 9, "Type": "DATA"}], "deleted_rows": []}
    app.session_state[key] = delta
    app.run(timeout=15)
    assert not app.exception and not app.error
    assert "report" not in app.session_state
    app.session_state[key] = delta
    app.button[0].click().run(timeout=15)
    assert not app.exception and not app.error
    anomalies = app.session_state["report"]["anomalies"]
    assert len(anomalies.loc[(anomalies["Fichier"] == "A") & (anomalies["Cellule"] == "A9")]) == 1


def test_structural_only_workbook_can_be_viewed_and_exported(monkeypatch):
    import streamlit as st
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    w = Workbook()
    w.active.append(["CIN", "DDR", "Note"])
    w.active.append(["Report", 12345, None])
    data = BytesIO()
    w.save(data)
    monkeypatch.setattr(st, "file_uploader", lambda *args, **kwargs: BytesIO(data.getvalue()))
    app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py")).run(timeout=15)
    assert not app.button[0].disabled
    app.button[0].click().run(timeout=15)
    assert not app.exception and not app.error
    assert app.session_state["report"]["main_metrics"]["Nombre de lignes analysées"] == 0
    assert app.session_state["export"]
