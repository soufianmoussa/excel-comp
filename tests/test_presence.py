from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_reader import discover_tables, read_table
from quality_report import missing_people


def workbook(rows):
    w = Workbook()
    w.active.append(["CIN", "DDR", "NOM ET PRENOM"])
    for row in rows:
        w.active.append(row)
    return w


def table(rows):
    w = workbook(rows)
    return read_table(w, discover_tables(w)[0])


def test_presence_both_directions_repetitions_and_structural_rows():
    a = table([["X1", 1, "Commun"], ["X2", 2, "Seulement A"], ["X2", 2, "Seulement A"], ["Report", 50, None]])
    b = table([["X3", 3, "Seulement B"], ["X1", 1, "Commun"]])
    result = missing_people(a, b, ["cin", "ddp", "name"])
    assert result["CIN"].tolist() == ["X2", "X2", "X3"]
    assert result["Ligne Excel"].tolist() == [3, 4, 2]
    assert result["Absent du fichier"].tolist() == ["B", "B", "A"]


def test_uncertain_incomplete_and_conflicting_matches_are_not_absent():
    a = table([["X1", 1, "Commun"], [None, None, "Incomplet"]])
    b = table([["X1", 2, "Commun"], ["X2", 1, "Autre"]])
    assert missing_people(a, b, ["cin", "ddp", "name"]).empty
    assert missing_people(a, b, []).empty
    # B's name fallback references A, although A has a stronger match elsewhere.
    a = table([["X1", 1, "Autre"], ["X2", 2, "Même nom"]])
    b = table([["X1", 1, "Même nom"]])
    assert missing_people(a, b, ["cin", "ddp", "name"]).empty


def test_presence_section_and_export(monkeypatch):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    data = {}
    for label, rows in (("A", [["X1", 1, "Commun"], ["X2", 2, "Seulement A"]]),
                        ("B", [["X1", 1, "Commun"], ["X3", 3, "Seulement B"]])):
        buf = BytesIO()
        workbook(rows).save(buf)
        data[f"upload_{label}"] = buf.getvalue()
    monkeypatch.setattr(st, "file_uploader", lambda *args, **kwargs: BytesIO(data[kwargs["key"]]))
    app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py")).run(timeout=15)
    app.button[0].click().run(timeout=15)
    assert not app.exception and not app.error
    assert any(s.value == "Personnes absentes d’un fichier" for s in app.subheader)
    assert any(t.label == "Présentes dans A, absentes de B (1)" for t in app.tabs)
    assert any(t.label == "Présentes dans B, absentes de A (1)" for t in app.tabs)
    assert len(app.metric) == 3
    exported = load_workbook(BytesIO(app.session_state["export"]))
    for title, cin in (("Dans A absentes de B", "X2"), ("Dans B absentes de A", "X3")):
        rows = list(exported[title].values)
        assert len(rows) == 2
        assert rows[1][rows[0].index("CIN")] == cin
