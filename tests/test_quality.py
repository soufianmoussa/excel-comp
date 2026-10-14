from copy import deepcopy
from datetime import datetime
from io import BytesIO
from time import perf_counter

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from comparator import compare
from data_quality import Rule, analyze_table, default_rule, frame_rules, rules_frame, parse_date, parse_number, validate_rule
from excel_reader import detect_tables, discover_tables, read_table
from exporter import export_quality
from field_mapping import detect_mappings, compare_mapped, sources
from quality_report import analyze, presentation_summary, errors_only_report
from quality_ui import styled_table


def book(headers, rows, title="Données"):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    table = read_table(workbook, discover_tables(workbook)[0])
    return workbook, table


def full_report(a, b, mappings=None, rules_a=None, rules_b=None):
    mappings = detect_mappings(a, b, ["cin", "ddp"]) if mappings is None else mappings
    return analyze([a], [b], a, b, ["cin", "ddp"], mappings, rules_a or {}, rules_b or {})


def test_all_columns_including_unnamed_duplicates_and_unknown_only_row():
    _, table = book(["CIN", "DDR", "Libellé", "Libellé", None],
                    [["X1", 1, "alpha", "beta", "extra"], [None, None, "alone", None, None]])
    assert len(table.all_columns) == 5
    assert table.rows == [2, 3]
    assert table.cells[2, 5].coordinate == "E2"
    assert table.cells[3, 3].display == "alone"


def test_generic_headers_and_all_sheets_inventory():
    w, _ = book(["CIN", "DDR"], [["X1", 1]])
    s = w.create_sheet("Autre")
    s.append(["Document"])
    s.append(["Référence", "Commentaire"])
    s.append([1, " suivi "])
    w.create_sheet("Vide")
    tables = [read_table(w, t) for t in discover_tables(w)]
    assert len(tables) == 2
    assert tables[1].header == 2
    assert set(tables[1].all_columns.values()) == {"Référence", "Commentaire"}
    assert "Espaces" in analyze_table(tables[1], "A")["Type"].tolist()


def test_required_optional_and_independent_unmatched_b():
    _, a = book(["CIN", "DDR", "Note"], [["X1", 1, None]])
    _, b = book(["CIN", "DDR", "Note"], [[None, 99, None]])
    report = full_report(a, b)
    missing = report["anomalies"].query("Type == 'Valeur manquante'")
    assert len(missing) == 1 and missing.iloc[0]["Fichier"] == "B"
    assert missing.iloc[0]["Cellule"] == "A2"
    assert report["stats"]["Total de cellules analysées"] == 6
    rule = Rule(required=True, kind="text")
    assert len(analyze_table(a, "A", {3: rule}).query("Type == 'Valeur manquante'")) == 1


def test_dates_invalid_chronology_and_ranges():
    _, t = book(["CIN", "DDR", "DEBUT", "FIN"], [["X1", 1, "31/02/2025", None], ["X2", 2, "03/10/2025", "01/10/2025"]])
    findings = analyze_table(t, "A")
    assert findings.query("Type == 'Date invalide'")["Cellule"].tolist() == ["C2"]
    assert set(findings.query("Type == 'Chronologie'")["Cellule"]) == {"C3", "D3"}
    findings = analyze_table(t, "A", {3: Rule(kind="date", date_min="2026-01-01")})
    assert "Plage de dates" in findings["Type"].tolist()
    assert parse_date(datetime(2025, 1, 1)) == parse_date("01/01/2025")


def test_rules_number_length_allowed_custom_and_unknown_safety():
    _, t = book(["CIN", "DDR", "Valeur"], [["X1", 1, "ABC"], ["X2", 2, "ab"]])
    assert analyze_table(t, "A").query("Gravité != 'Information'").empty  # unfamiliar values aren't invented errors
    rule = Rule(kind="text", min_length=3, max_length=4, allowed="ABC|DEF", custom="prefixe:A")
    f = analyze_table(t, "A", {3: rule})
    assert set(f.query("Gravité == 'Erreur'")["Type"]) == {"Longueur", "Valeur autorisée", "Règle personnalisée"}
    _, t = book(["CIN", "DDR", "NET(DH)"], [["X1", 1, -2], ["X2", 2, "invalide"]])
    f = analyze_table(t, "A")
    assert "Plage numérique" not in f["Type"].tolist()
    f = analyze_table(t, "A", {3: Rule(kind="number", minimum="0", maximum="10")})
    assert set(f["Type"]) == {"Type", "Plage numérique"}
    assert parse_number("1 234,50") == parse_number("1234.5")


def test_rule_validation_rejects_code_and_invalid_bounds():
    for rule in [Rule(custom="eval:__import__('os')"), Rule(minimum="ten"), Rule(min_length=5, max_length=2),
                 Rule(minimum="10", maximum="2"), Rule(date_min="31/02/2025")]:
        with pytest.raises(ValueError):
            validate_rule(rule)


def test_capitalization_spaces_duplicate_rules_only():
    _, t = book(["CIN", "DDR", "CATEGORIE"], [["X1", 1, " Cadre "], ["X1", 2, "CADRE"]])
    f = analyze_table(t, "A")
    assert f.query("Type == 'Doublon'").empty
    assert len(f.query("Type == 'Espaces'")) == 1
    assert len(f.query("Type == 'Casse / écriture'")) == 2
    assert not f.query("Titre == 'CATEGORIE' and Gravité == 'Erreur'").shape[0]
    assert analyze_table(t, "A", {1: Rule(kind="identifier")}).query("Type == 'Doublon'").empty


def test_uncertain_numeric_and_formula_not_false_differences():
    _, a = book(["CIN", "DDR", "RIB", "Montant"], [["X1", 1, 123456789012345678901234, "=2+2"]])
    _, b = book(["CIN", "DDR", "RIB", "Montant"], [["X1", 1, "123456789012345678901234", 4]])
    mappings = detect_mappings(a, b, ["cin", "ddp", "rib"])
    r = full_report(a, b, mappings)
    assert "Précision numérique" in r["anomalies"]["Type"].tolist()
    assert "Formule non évaluée" in r["anomalies"]["Type"].tolist()
    assert r["differences"].query("Type == 'Différence'").empty
    assert len(r["differences"].query("Type == 'Comparaison impossible'")) == 2


def test_dates_mapped_with_both_cell_coordinates_and_coloring():
    _, a = book(["CIN", "DDR", "DEBUT"], [["X1", 1, "01/10/2025"]])
    _, b = book(["CIN", "DDR", "Date début"], [["X1", 1, "02/10/2025"]])
    r = full_report(a, b)
    assert r["records"].iloc[0]["Code statut"] == "DIFFERENT"
    diff = r["differences"].iloc[0]
    assert diff["Cellules A"] == diff["Cellules B"] == "C2"
    assert diff["Valeur A"] == "01/10/2025" and diff["Valeur B"] == "02/10/2025"
    assert r["stats"]["Différences entre fichiers"] == 1
    assert len(r["findings"].query("Type == 'Différence'")) == 2
    styled = styled_table(a, a.rows, r["findings"].query("Fichier == 'A'"))
    assert "FADBD8" in styled.to_html()


def test_no_fuzzy_mapping_and_net_requires_confirmation():
    _, a = book(["CIN", "DDR", "NET(DH)", "Prime A"], [["X1", 1, 1, 4]])
    _, b = book(["CIN", "DDR", "Montant net", "Prime B"], [["X1", 1, 2, 5]])
    m = detect_mappings(a, b, ["cin", "ddp"])
    assert not bool(m.loc[m["Champ"] == "Montant net", "Actif"].iloc[0])
    r = full_report(a, b, m)
    assert r["differences"].empty
    assert len(r["unmatched"]) == 4
    m.loc[m["Champ"] == "Montant net", "Actif"] = True
    assert len(full_report(a, b, m)["differences"]) == 1


def test_disable_mapping_and_add_custom_mapping():
    _, a = book(["CIN", "DDR", "État A"], [["X1", 1, "Oui"]])
    _, b = book(["CIN", "DDR", "État B"], [["X1", 2, "Non"]])
    m = detect_mappings(a, b, ["cin", "ddp"])
    m.loc[m["Mode"] == "ddp", "Actif"] = False
    r = full_report(a, b, m)
    assert r["records"].iloc[0]["Code statut"] == "MATCH"
    assert r["records"].iloc[0]["Contrôle DDP / DDR"] == "Non comparé"
    m.loc[len(m)] = [True, "État", list(sources(a))[2], list(sources(b))[2], "text", "Administrateur"]
    r = full_report(a, b, m)
    assert r["records"].iloc[0]["Code statut"] == "DIFFERENT"
    assert r["differences"].iloc[0]["Champ"] == "État"


def test_no_active_mapping_never_claims_complete_match():
    _, a = book(["CIN", "DDR"], [["X1", 1]])
    m = detect_mappings(a, a, ["cin", "ddp"])
    m["Actif"] = False
    assert full_report(a, a, m)["records"].iloc[0]["Code statut"] == "INCOMPLETE"


def test_optional_selected_identifier_does_not_create_incomplete_match():
    _, a = book(["CIN", "DDR"], [[None, 1]])
    rule = {(a.sheet, 1): Rule(kind="identifier", required=False)}
    report = full_report(a, a, rules_a=rule, rules_b=rule)
    assert report["records"].iloc[0]["Code statut"] == "MATCH"
    assert report["anomalies"].empty
    required_report = full_report(a, a)
    assert required_report["records"].iloc[0]["Code statut"] == "INCOMPLETE"


def test_simple_summary_counts_both_files_repeated_rows_and_additional_sheets():
    _, a = book(["CIN", "DDR", "Note"], [["X1", 1, None], ["X1", 1, None]])
    _, b = book(["CIN", "DDR", "Note"], [["X1", 1, None]])
    _, extra = book(["Commentaire", "Valeur"], [["texte", 3]], title="Annexe")
    m = detect_mappings(a, b, ["cin", "ddp"])
    report = analyze([a, extra], [b], a, b, ["cin", "ddp"], m, {}, {})
    items, metrics = presentation_summary(report, [a, extra], [b], a, b)
    assert list(metrics.values()) == [4, 0, 4]
    assert items.empty


def test_simple_summary_ignores_uncertainty_but_keeps_matching_status():
    _, a = book(["CIN", "DDR"], [["X1", 1], ["X9", 9]])
    _, b = book(["CIN", "DDR"], [["X1", 1], ["X1", 1]])
    report = full_report(a, b)
    items, metrics = presentation_summary(report, [a], [b], a, b)
    assert list(metrics.values()) == [4, 0, 4]
    assert items.empty
    assert report["records"]["Code statut"].tolist() == ["AMBIGUOUS", "MISSING_IN_B"]
    assert not any("Doublon" in t for t in items["Type"])


def test_summary_does_not_count_same_problematic_row_twice():
    _, a = book(["CIN", "DDR", "DEBUT", "FIN"], [["X1", 1, "31/02/2025", "31/02/2025"]])
    _, b = book(["CIN", "DDR", "DEBUT", "FIN"], [["X1", 1, "01/02/2025", "02/02/2025"]])
    report = full_report(a, b)
    items, metrics = presentation_summary(report, [a], [b], a, b)
    assert list(metrics.values()) == [2, len(items), 1]
    assert len(items) == 2


def test_only_errors_survive_in_details_counts_and_export():
    wa, a = book(["CIN", "DDR", "DEBUT", "Note", "Montant"], [["X1", 1, "31/02/2025", " texte ", "=1+1"]])
    wb, b = book(["CIN", "DDR", "DEBUT", "Note", "Montant"], [["X1", 1, "01/02/2025", "texte", 2]])
    mappings = detect_mappings(a, b, ["cin", "ddp"])
    raw = full_report(a, b, mappings)
    assert set(raw["findings"]["Gravité"]) == {"Erreur", "Information", "Avertissement"}
    report = errors_only_report(raw)
    for key in ("findings", "anomalies", "differences"):
        assert set(report[key]["Gravité"]) <= {"Erreur"}
    assert report["stats"]["Anomalies du fichier A"] == 1
    assert report["stats"]["Anomalies du fichier B"] == 0
    assert "Comparaisons impossibles" not in report["stats"]
    # Raw reports are filtered defensively by the exporter too.
    out = load_workbook(BytesIO(export_quality(raw, wa, wb, [a], [b], a, b, mappings, rules_frame([a]), rules_frame([b]))))
    sheet = out["Fichier A annoté"]
    assert sheet["C2"].comment is not None
    assert sheet["D2"].comment is None and sheet["D2"].fill.patternType is None
    assert sheet["E2"].comment is None and sheet["E2"].fill.patternType is None
    for title in ("Anomalies fichier A", "Anomalies fichier B", "Différences entre fichiers"):
        values = list(out[title].values)
        severity = list(values[0]).index("Gravité")
        assert all(row[severity] == "Erreur" for row in values[1:])


def test_unknown_codes_with_leading_zeros_not_compared_as_numbers():
    _, a = book(["CIN", "DDR", "Code local"], [["X1", 1, "001"]])
    _, b = book(["CIN", "DDR", "Code local"], [["X1", 1, "1"]])
    report = full_report(a, b)
    assert len(report["differences"]) == 1


def test_merged_coordinates_and_styled_export():
    wa, a = book(["CIN", "DDR", "Note"], [["X1", 1, " info "], ["X2", 2, None]])
    wa.active.merge_cells("C2:C3")
    a = read_table(wa, discover_tables(wa)[0])
    wb, b = book(["CIN", "DDR", "Note"], [["X1", 1, "other"], ["X2", 2, "other"]])
    m = detect_mappings(a, b, ["cin", "ddp"])
    report = full_report(a, b, m)
    assert a.cells[3, 3].coordinate == "C3"
    assert a.cells[3, 3].value is None
    assert a.cells[3, 3].role == "MERGED_PLACEHOLDER"
    out = load_workbook(BytesIO(export_quality(report, wa, wb, [a], [b], a, b, m, rules_frame([a]), rules_frame([b]))))
    assert "C2:C3" in out["Fichier A annoté"].merged_cells
    assert out["Fichier A annoté"]["C2"].comment is not None


def test_composed_sources_preserve_all_coordinates():
    rib = "001002000000000000000304"
    _, a = book(["CIN", "DDR", "NOM ET PRENOM", "RIB"], [["X1", 1, "Said Exemple", rib]])
    _, b = book(["CIN", "DDR", "Prénom", "Nom", "Banque", "Ville", "Numéro compte RIB", "Clé"],
                [["X1", 1, "Said", "Exemple", 1, 2, "0000000000000003", 5]])
    m = detect_mappings(a, b, ["cin", "ddp", "name", "rib"])
    r = full_report(a, b, m)
    assert len(r["differences"]) == 1
    assert r["differences"].iloc[0]["Cellules B"] == "E2, F2, G2, H2"


def test_annotated_export_preserves_format_original_values_and_coordinates():
    wa, a = book(["CIN", "DDR", "DEBUT", "Note"], [["X1", 1, "01/10/2025", "=HYPERLINK(\"external\")"]])
    wb, b = book(["CIN", "DDR", "Date début", "Note"], [["X1", 1, "02/10/2025", None]])
    wa.active["C2"].font = Font(bold=True, color="123456")
    wa.active.column_dimensions["C"].width = 27
    m = detect_mappings(a, b, ["cin", "ddp"])
    r = full_report(a, b, m)
    data = export_quality(r, wa, wb, [a], [b], a, b, m, rules_frame([a]), rules_frame([b]))
    out = load_workbook(BytesIO(data))
    assert {"Résumé global", "Anomalies fichier A", "Anomalies fichier B", "Différences entre fichiers", "Correspondances ambiguës", "Colonnes non comparées", "Fichier A annoté", "Fichier B annoté"} <= set(out.sheetnames)
    cell = out["Fichier A annoté"]["C2"]
    assert cell.value == "01/10/2025" and cell.font.bold and cell.font.color.rgb == "00123456"
    assert cell.fill.fgColor.rgb == "00FADBD8" and "Différence" in cell.comment.text
    assert out["Fichier A annoté"].column_dimensions["C"].width == 27
    assert out["Fichier A annoté"]["D2"].data_type == "s"
    assert wa.active["C2"].comment is None  # original never mutated


def test_rules_roundtrip_and_inferred_types_are_warnings():
    _, t = book(["CIN", "DDR", "Mesure"], [[f"X{i}", i, i if i != 9 else "abc"] for i in range(10)])
    f = rules_frame([t])
    rules = frame_rules(f)
    assert rules[t.sheet, 3].kind == "auto"
    findings = analyze_table(t, "A")
    wrong_type = findings.query("Titre == 'Mesure' and Type == 'Type'")
    assert len(wrong_type) == 1 and wrong_type.iloc[0]["Gravité"] == "Avertissement"


def test_thousand_rows_all_columns_complete_without_cartesian_match():
    headers = ["CIN", "DDR", "Date début", *[f"Mesure {i}" for i in range(8)]]
    rows = [[f"X{i}", i, "01/10/2025", *range(8)] for i in range(1500)]
    start = perf_counter()
    _, a = book(headers, rows)
    _, b = book(headers, list(reversed(rows)))
    report = full_report(a, b)
    assert report["stats"]["Total de cellules analysées"] == 33000
    assert report["differences"].empty
    assert (report["records"]["Code statut"] == "MATCH").all()
    # Generous regression guard; ordinary execution is substantially faster.
    assert perf_counter() - start < 30
