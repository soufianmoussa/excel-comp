"""French configuration controls and paginated spreadsheet viewers."""
import hashlib
import math

import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter

from data_quality import frame_rules, rules_frame, SEVERITIES
from exporter import COLORS
from field_mapping import detect_mappings, sources, validate_mappings


ROW_LABELS = {"DATA": "Données", "HEADER": "En-têtes", "TITLE": "Titre", "REPORT": "Report",
              "TOTAL": "Total", "SUBTOTAL": "Sous-total", "SUMMARY": "Synthèse", "NOTES": "Notes",
              "EMPTY": "Vide", "UNKNOWN": "Structure à vérifier", "OTHER": "Structure à vérifier"}


def configure(tables_a, tables_b, a, b, fields, key):
    valid = True
    with st.expander("Configuration avancée"):
        st.caption("Réservé aux personnes qui connaissent les règles métier. Les colonnes inconnues sont facultatives et les déductions automatiques restent des avertissements.")
        st.markdown("**Correspondances entre colonnes**")
        st.caption("Les intitulés semblables ne sont pas rapprochés automatiquement. Confirmez le sens métier avant d'activer une proposition. Ajoutez ou supprimez des lignes pour corriger la configuration. Une correspondance inactive ne participe pas au contrôle détaillé.")
        field_key = "_".join(fields)
        mapping = st.data_editor(detect_mappings(a, b, fields), key=f"mappings_{key}_{field_key}", num_rows="dynamic", hide_index=True, width="stretch",
            column_config={"Actif": st.column_config.CheckboxColumn("Comparer", default=False),
                           "Source A": st.column_config.SelectboxColumn("Colonne fichier A", options=list(sources(a)), required=True),
                           "Source B": st.column_config.SelectboxColumn("Colonne fichier B", options=list(sources(b)), required=True),
                           "Mode": st.column_config.SelectboxColumn("Type de comparaison", options=["name", "ddp", "cin", "rib", "text", "number", "date", "identifier"], required=True)},
            disabled=["Origine"])
        st.caption("Modes : name = nom sans casse ni accents ; ddp / cin / rib = normalisation des identifiants ; text = texte ; number = nombre ; date = date ; identifier = identifiant.")
        edited, parsed = {}, {}
        for label, tables in (("A", tables_a), ("B", tables_b)):
            st.markdown(f"**Règles du fichier {label} — toutes les colonnes**")
            frame = rules_frame(tables)
            for column in ("min_length", "max_length"):
                frame[column] = pd.to_numeric(frame[column], errors="coerce").astype(float)
            edited[label] = st.data_editor(frame, key=f"rules_{label}_{key}", hide_index=True, width="stretch",
                disabled=["Feuille", "Colonne", "Titre", "Type suggéré"],
                column_config={
                    "required": st.column_config.CheckboxColumn("Obligatoire"),
                    "kind": st.column_config.SelectboxColumn("Type attendu", options=["auto", "text", "number", "date", "identifier"], required=True),
                    "min_length": st.column_config.NumberColumn("Longueur min.", min_value=0, step=1),
                    "max_length": st.column_config.NumberColumn("Longueur max.", min_value=0, step=1),
                    "allowed": st.column_config.TextColumn("Valeurs autorisées", help="Séparer les valeurs par | ; vide = aucune restriction."),
                    "minimum": st.column_config.TextColumn("Nombre min."), "maximum": st.column_config.TextColumn("Nombre max."),
                    "date_min": st.column_config.TextColumn("Date min.", help="JJ/MM/AAAA"),
                    "date_max": st.column_config.TextColumn("Date max.", help="JJ/MM/AAAA"),
                    "custom": st.column_config.TextColumn("Règle personnalisée", help="prefixe:ABC, suffixe:XYZ, contient:texte ou sans:texte"),
                })
        st.caption("Laissez les bornes vides pour ne pas les imposer. Les dates début/fin sont facultatives par défaut. Aucun montant négatif n'est interdit sans minimum configuré. Règle personnalisée : prefixe:…, suffixe:…, contient:… ou sans:… (texte exact, sans code exécutable).")
        try:
            parsed = {label: frame_rules(frame) for label, frame in edited.items()}
            validate_mappings(mapping, a, b)
        except (ValueError, TypeError) as error:
            st.error(str(error))
            valid = False
    return mapping, edited, parsed, valid


def table_frame(table, rows=None):
    rows = (table.original_rows or table.rows) if rows is None else rows
    frame = pd.DataFrame([[table.cells[row, col].display for col in table.all_columns] for row in rows],
        index=pd.Index(rows, name="Ligne Excel"),
        columns=[f"{get_column_letter(col)} · {title}" for col, title in table.all_columns.items()])
    frame.insert(0, "Type de ligne", [ROW_LABELS.get(table.row_classifications[r].kind, table.row_classifications[r].kind)
                                    if r in table.row_classifications else "Données" for r in rows])
    return frame


def styled_table(table, rows, findings):
    frame = table_frame(table, rows)
    ranks = {severity: index for index, severity in enumerate(SEVERITIES)}
    marked = {}
    for item in findings.to_dict("records"):
        if item["Gravité"] != "Erreur":
            continue
        key = item["Ligne"], item["Colonne"]
        if key not in marked or ranks[item["Gravité"]] < ranks[marked[key]]:
            marked[key] = item["Gravité"]
    styles = pd.DataFrame("", index=frame.index, columns=frame.columns)
    for (row, col), severity in marked.items():
        if row in frame.index and col in table.all_columns:
            column = f"{get_column_letter(col)} · {table.all_columns[col]}"
            styles.at[row, column] = f"background-color: #{COLORS[severity]}; color: #172E32"
    return frame.style.apply(lambda _: styles, axis=None)


def missing_people_view(report, a, b):
    st.subheader("Personnes absentes d’un fichier")
    st.caption(f"Feuilles comparées : A — {a.sheet} ; B — {b.sheet}. Les listes concernent les personnes sans correspondance avec les identifiants disponibles. Les identités incomplètes ou les correspondances incertaines ne sont pas déclarées absentes. Les lignes répétées sont conservées.")
    frame = report["missing_people"]
    only_a = frame.loc[frame["Fichier"] == "A"]
    only_b = frame.loc[frame["Fichier"] == "B"]
    tabs = st.tabs([f"Présentes dans A, absentes de B ({len(only_a)})",
                    f"Présentes dans B, absentes de A ({len(only_b)})"])
    for tab, subset in zip(tabs, (only_a, only_b)):
        with tab:
            if subset.empty:
                st.caption("Aucune absence établie dans ce sens.")
            else:
                st.dataframe(subset.drop(columns=["Fichier", "Absent du fichier"]), hide_index=True, width="stretch")


def viewer(report, tables_a, tables_b, a, b):
    with st.expander("Structures à vérifier"):
        uncertain = [{"Fichier": label, "Feuille": t.sheet, "Ligne Excel": r,
                      "Confiance": info.confidence, "Raison": info.reason}
                     for label, tables in (("A", tables_a), ("B", tables_b)) for t in tables
                     for r, info in t.row_classifications.items() if info.kind in {"UNKNOWN", "OTHER"}]
        st.caption("La classification est heuristique. Les zones incertaines sont conservées sans contrôle de bénéficiaire. Corrigez leur type dans les options avancées du fichier, puis relancez l’analyse.")
        if uncertain:
            st.dataframe(pd.DataFrame(uncertain), hide_index=True, width="stretch")
        else:
            st.caption("Aucune structure incertaine détectée.")
    st.subheader("Erreurs détectées")
    st.caption("Cliquez sur une erreur pour afficher la cellule concernée en rouge. Une différence ne désigne pas le fichier correct.")
    findings = report.get("visible_findings", report["findings"])
    findings = findings.loc[findings["Gravité"] == "Erreur"]
    controls = st.columns(3)
    file_filter = controls[0].selectbox("Fichier", ["A et B", "A", "B"])
    type_filter = controls[1].multiselect("Type d'anomalie", sorted(findings["Type"].unique()))
    columns = sorted(findings["Titre"].unique())
    # Include clean columns too, so every source column can be filtered.
    columns = sorted(set(columns) | {title for t in [*tables_a, *tables_b] for title in t.all_columns.values()})
    column_filter = controls[2].multiselect("Colonne", columns)
    search = st.text_input("Rechercher une personne ou une valeur", placeholder="Nom, identifiant, date ou autre valeur…")
    only_errors = st.checkbox("Seulement les lignes avec anomalies", value=False)
    filtered = findings.copy()
    if file_filter != "A et B":
        filtered = filtered.loc[filtered["Fichier"] == file_filter]
    if type_filter:
        filtered = filtered.loc[filtered["Type"].isin(type_filter)]
    if column_filter:
        filtered = filtered.loc[filtered["Titre"].isin(column_filter)]
    if search:
        # A search matches entire source rows, including values in another column.
        matched_rows = set()
        for label, tables in (("A", tables_a), ("B", tables_b)):
            for table in tables:
                for row in (table.original_rows or table.rows):
                    if any(search.casefold() in table.cells[row, col].display.casefold() for col in table.all_columns):
                        matched_rows.add((label, table.sheet, row))
        filtered = filtered.loc[[tuple(x) in matched_rows for x in filtered[["Fichier", "Feuille", "Ligne"]].itertuples(index=False, name=None)]]
    focus = None
    if not filtered.empty:
        filtered = filtered.reset_index(drop=True)
        nav_key = "anomaly_" + hashlib.sha256(filtered.to_json().encode()).hexdigest()[:12]
        if nav_key not in st.session_state:
            st.session_state[nav_key] = 0
        visible_columns = ["Fichier", "Feuille", "Ligne", "Personne", "Titre", "Type", "Explication"]
        simple = filtered[[c for c in visible_columns if c in filtered]].rename(columns={"Ligne": "Ligne Excel", "Titre": "Colonne", "Type": "Type d’anomalie", "Explication": "Description du problème"})
        event = st.dataframe(simple, hide_index=True, width="stretch", key=f"list_{nav_key}",
                             on_select="rerun", selection_mode="single-row")
        selected_rows = event.selection.rows
        selection_key = f"last_selection_{nav_key}"
        if selected_rows and 0 <= selected_rows[0] < len(filtered) and selected_rows != st.session_state.get(selection_key):
            st.session_state[nav_key] = selected_rows[0]
        st.session_state[selection_key] = selected_rows
        nav = st.columns([1, 5, 1])
        def move(delta):
            st.session_state[nav_key] = (st.session_state[nav_key] + delta) % len(filtered)
        nav[0].button("← Précédente", on_click=move, args=(-1,), key=f"previous_{nav_key}")
        selected = nav[1].selectbox("Naviguer entre les anomalies", list(range(len(filtered))),
            format_func=lambda i: f"{i + 1}/{len(filtered)} · {filtered.iloc[i]['Fichier']} / {filtered.iloc[i]['Feuille']} / {filtered.iloc[i]['Cellule']} · {filtered.iloc[i]['Type']}", key=nav_key)
        nav[2].button("Suivante →", on_click=move, args=(1,), key=f"next_{nav_key}")
        focus = filtered.iloc[selected]
        st.caption(f"Règle : {focus.get('Règle appliquée', focus['Type'])}. Contexte : {focus.get('Contexte de validation', '')}")
        st.info(f"{focus['Gravité']} — {focus['Explication']} (fichier {focus['Fichier']}, feuille {focus['Feuille']}, cellule {focus['Cellule']})")
        st.caption("La cellule sélectionnée est visible dans l'extrait ci-dessous. Le tableau complet reste consultable plus bas.")
        table = next(t for t in (tables_a if focus["Fichier"] == "A" else tables_b) if t.sheet == focus["Feuille"])
        marks = findings.loc[(findings["Fichier"] == focus["Fichier"]) & (findings["Feuille"] == table.sheet)]
        focus_column = f"{get_column_letter(int(focus['Colonne']))} · {table.all_columns[int(focus['Colonne'])]}"
        other_columns = [c for c in table_frame(table, []).columns if c != focus_column]
        st.dataframe(styled_table(table, [int(focus["Ligne"])], marks), column_order=[focus_column, *other_columns], width="stretch")
        differences = report["differences"]
        matching = differences.loc[(differences[f"Feuille {focus['Fichier']}"] == focus["Feuille"]) &
                                   (differences[f"Ligne {focus['Fichier']}"] == focus["Ligne"])]
        if not matching.empty:
            st.caption("Valeurs correspondantes dans les deux fichiers")
            st.dataframe(matching[["Champ", "Cellules A", "Valeur A", "Cellules B", "Valeur B"]], hide_index=True, width="stretch")
    else:
        st.info("Aucune anomalie pour ces filtres.")
    st.subheader("Explorer les fichiers")
    for container, label, tables in zip(st.columns(2), ("A", "B"), (tables_a, tables_b)):
        if file_filter not in {"A et B", label}:
            continue
        with container:
            st.markdown(f"**Fichier {label} — toutes les colonnes**")
            sheet = st.selectbox("Feuille à afficher", [t.sheet for t in tables], key=f"view_sheet_{label}")
            table = next(t for t in tables if t.sheet == sheet)
            frame = table_frame(table)
            if search:
                frame = frame.loc[frame.apply(lambda col: col.str.contains(search, case=False, regex=False)).any(axis=1)]
            selected_marks = filtered.loc[(filtered["Fichier"] == label) & (filtered["Feuille"] == sheet)]
            if only_errors or type_filter:
                frame = frame.loc[frame.index.isin(selected_marks["Ligne"])]
            pages = max(1, math.ceil(len(frame) / 100))
            page = st.number_input("Page (100 lignes)", min_value=1, max_value=pages, value=1,
                                   key=f"page_{label}_{sheet}_{len(frame)}")
            rows = frame.index[(int(page) - 1) * 100:int(page) * 100].tolist()
            marks = findings.loc[(findings["Fichier"] == label) & (findings["Feuille"] == sheet)]
            st.caption(f"{len(frame)} lignes · {len(table.all_columns)} colonnes · références Excel conservées")
            view = styled_table(table, rows, marks)
            if column_filter:
                visible = [f"{get_column_letter(c)} · {title}" for c, title in table.all_columns.items() if title in column_filter]
                st.dataframe(view, column_order=visible, width="stretch", height=400)
            else:
                st.dataframe(view, width="stretch", height=400)
    records = report["records"]
    if not records.empty:
        pairs = records.loc[records["Lignes B"].astype(str).str.fullmatch(r"\d+") & (records["Code statut"] != "AMBIGUOUS")]
        if not pairs.empty:
            st.markdown("**Comparer une personne côte à côte**")
            selected = st.selectbox("Lignes correspondantes", pairs.index.tolist(),
                format_func=lambda i: f"A : {a.sheet}, ligne {pairs.at[i, 'Ligne A']} ↔ B : {b.sheet}, ligne {pairs.at[i, 'Lignes B']}")
            for container, label, table, row in zip(st.columns(2), ("A", "B"), (a, b),
                    (int(pairs.at[selected, "Ligne A"]), int(pairs.at[selected, "Lignes B"]))):
                with container:
                    st.caption(f"Fichier {label} · {table.sheet} · ligne {row}")
                    marks = findings.loc[(findings["Fichier"] == label) & (findings["Feuille"] == table.sheet)]
                    st.dataframe(styled_table(table, [row], marks), width="stretch")
    with st.expander("Liste détaillée des anomalies filtrées"):
        st.dataframe(filtered, hide_index=True, width="stretch")
